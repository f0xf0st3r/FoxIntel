"""Office document analyzer"""

from typing import Dict
from pathlib import Path
from core.base import BaseModule
from core.exceptions import ValidationError


class OfficeAnalyzer(BaseModule):
    def analyze(self, file_path: str) -> Dict:
        if not Path(file_path).exists():
            raise ValidationError(f"File not found: {file_path}")

        self._log_info(f"Analyzing Office document: {file_path}")

        results = {
            "file": file_path,
            "type": self._detect_type(file_path),
            "metadata": {},
            "content": {},
            "macros": {},
            "links": []
        }

        ext = Path(file_path).suffix.lower()

        if ext in [".doc", ".docx"]:
            results.update(self._analyze_word(file_path))
        elif ext in [".xls", ".xlsx"]:
            results.update(self._analyze_excel(file_path))
        elif ext in [".ppt", ".pptx"]:
            results.update(self._analyze_powerpoint(file_path))

        return self._format_result(True, results)

    def _detect_type(self, file_path: str) -> str:
        ext = Path(file_path).suffix.lower()
        types = {
            ".doc": "word_document",
            ".docx": "word_document",
            ".xls": "excel_spreadsheet",
            ".xlsx": "excel_spreadsheet",
            ".ppt": "powerpoint_presentation",
            ".pptx": "powerpoint_presentation"
        }
        return types.get(ext, "unknown")

    def _analyze_word(self, file_path: str) -> Dict:
        results = {}
        try:
            from docx import Document

            doc = Document(file_path)

            results["metadata"] = {
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "sections": len(doc.sections)
            }

            results["content"]["text"] = "\n".join([p.text for p in doc.paragraphs])

            core_props = doc.core_properties
            results["metadata"]["author"] = core_props.author
            results["metadata"]["created"] = str(core_props.created)
            results["metadata"]["modified"] = str(core_props.modified)

        except ImportError:
            results["error"] = "python-docx not installed"
        except Exception as e:
            results["error"] = str(e)

        return results

    def _analyze_excel(self, file_path: str) -> Dict:
        results = {}
        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path, data_only=True)

            results["metadata"] = {
                "sheets": len(wb.sheetnames),
                "sheet_names": wb.sheetnames
            }

            results["content"]["sheets"] = {}
            for sheet_name in wb.sheetnames[:5]:
                sheet = wb[sheet_name]
                results["content"]["sheets"][sheet_name] = {
                    "rows": sheet.max_row,
                    "columns": sheet.max_column
                }

        except ImportError:
            results["error"] = "openpyxl not installed"
        except Exception as e:
            results["error"] = str(e)

        return results

    def _analyze_powerpoint(self, file_path: str) -> Dict:
        results = {}
        try:
            from pptx import Presentation

            prs = Presentation(file_path)

            results["metadata"] = {
                "slides": len(prs.slides),
                "title": prs.core_properties.title,
                "author": prs.core_properties.author
            }

            results["content"]["slide_count"] = len(prs.slides)

        except ImportError:
            results["error"] = "python-pptx not installed"
        except Exception as e:
            results["error"] = str(e)

        return results